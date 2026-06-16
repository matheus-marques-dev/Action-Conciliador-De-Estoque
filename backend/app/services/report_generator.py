import io
import uuid
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models.schemas import DivergenceStatus, ReconciliationItem

# ── Colour palette ──────────────────────────────────────────────────────────
_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_STATUS_FILL: Dict[DivergenceStatus, str] = {
    DivergenceStatus.ZERO_STOCK: "FF4444",
    DivergenceStatus.NOT_FOUND:  "FF8C69",
    DivergenceStatus.LOW_STOCK:  "FFB347",
    DivergenceStatus.OVERSTOCK:  "FFD700",
    DivergenceStatus.OK:         "90EE90",
}
_STATUS_LABEL: Dict[DivergenceStatus, str] = {
    DivergenceStatus.ZERO_STOCK: "Estoque Zero",
    DivergenceStatus.NOT_FOUND:  "Não Encontrado",
    DivergenceStatus.LOW_STOCK:  "Estoque Baixo",
    DivergenceStatus.OVERSTOCK:  "Excesso de Estoque",
    DivergenceStatus.OK:         "OK",
}

_THIN = Side(border_style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

COLUMNS = [
    ("SKU",              14),
    ("Descrição",        36),
    ("Marketplace",      18),
    ("ID Anúncio",       20),
    ("Qtd ERP",          10),
    ("Qtd Marketplace",  16),
    ("Divergência",      12),
    ("Status",           18),
    ("Ação Sugerida",    52),
]


def _header_row(ws, row: int = 1) -> None:
    for col, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 28


def _write_items(ws, items: List[ReconciliationItem], start_row: int = 2) -> None:
    for r, item in enumerate(items, start=start_row):
        status_color = _STATUS_FILL[item.status]
        fill = PatternFill("solid", fgColor=status_color)

        values = [
            item.sku,
            item.description,
            item.marketplace,
            item.listing_id or "—",
            item.erp_qty,
            item.marketplace_qty if item.marketplace_qty is not None else "—",
            item.divergence if item.divergence is not None else "—",
            _STATUS_LABEL[item.status],
            item.suggested_action,
        ]

        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c == len(values)))
            if c == 8:  # Status column gets colour
                cell.fill = fill
                cell.font = Font(bold=True)

        ws.row_dimensions[r].height = 18


def _summary_sheet(ws, items: List[ReconciliationItem]) -> None:
    from collections import Counter
    ws.title = "Resumo"

    counts = Counter(i.status for i in items)
    total = len(items)

    stats = [
        ("Total de SKUs analisados", total, "4A90D9"),
        ("OK — estoques alinhados",  counts[DivergenceStatus.OK], "90EE90"),
        ("Estoque Baixo",            counts[DivergenceStatus.LOW_STOCK], "FFB347"),
        ("Excesso de Estoque",       counts[DivergenceStatus.OVERSTOCK], "FFD700"),
        ("Estoque Zero (crítico)",   counts[DivergenceStatus.ZERO_STOCK], "FF4444"),
        ("Não Encontrado",           counts[DivergenceStatus.NOT_FOUND], "FF8C69"),
    ]

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12

    title_cell = ws.cell(row=1, column=1, value="Conciliação ERP × Marketplace")
    title_cell.font = Font(bold=True, size=14, color="1E3A5F")
    ws.merge_cells("A1:B1")

    for row, (label, count, color) in enumerate(stats, start=3):
        lc = ws.cell(row=row, column=1, value=label)
        lc.alignment = Alignment(vertical="center")
        lc.border = _BORDER
        vc = ws.cell(row=row, column=2, value=count)
        vc.fill = PatternFill("solid", fgColor=color)
        vc.font = Font(bold=True)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border = _BORDER
        ws.row_dimensions[row].height = 22


def generate_report(items: List[ReconciliationItem]) -> Tuple[str, bytes]:
    report_id = uuid.uuid4().hex[:8].upper()
    wb = openpyxl.Workbook()

    # Sheet 1 — Summary
    ws_summary = wb.active
    _summary_sheet(ws_summary, items)

    # Sheet 2 — Divergences only
    ws_div = wb.create_sheet("Divergências")
    ws_div.title = "Divergências"
    _header_row(ws_div)
    divergences = [i for i in items if i.status != DivergenceStatus.OK]
    _write_items(ws_div, divergences)
    ws_div.freeze_panes = "A2"

    # Sheet 3 — Full comparison
    ws_full = wb.create_sheet("Comparação Completa")
    _header_row(ws_full)
    _write_items(ws_full, items)
    ws_full.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return report_id, buf.read()
